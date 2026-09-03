"""Generate the .xlsx view of the datastore.

The spreadsheet is disposable: it is rebuilt from the datastore on every run,
with a frozen bold header row, auto-sized columns, and clickable hyperlinks in
the post-link column.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .store import StoredPost

# (header, attribute on StoredPost). Column order is the deliverable's order,
# with the two cheap extras (search term, capture time) after it.
COLUMNS = [
    ("Post Link", "url"),
    ("Date Posted", "date_posted"),
    ("Posted By", "author"),
    ("Post Text", "text"),
    ("Search Term", "search_term"),
    ("Captured At", "captured_at"),
]

LINK_COLUMN_INDEX = 1  # "Post Link" is column A.

# Upper bounds so one very long post doesn't create an unreadably wide column.
MAX_COL_WIDTH = 80
MIN_COL_WIDTH = 12


def write_xlsx(posts: list[StoredPost], path: str | Path) -> Path:
    """Write ``posts`` to an .xlsx file at ``path`` and return the path."""
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "LinkedIn Posts"

    # Header row.
    header_font = Font(bold=True)
    for col_idx, (header, _attr) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font

    # Data rows.
    hyperlink_font = Font(color="0563C1", underline="single")
    for row_idx, post in enumerate(posts, start=2):
        for col_idx, (_header, attr) in enumerate(COLUMNS, start=1):
            value = getattr(post, attr)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx == LINK_COLUMN_INDEX and value:
                cell.hyperlink = value
                cell.font = hyperlink_font

    _autosize_columns(ws, posts)

    # Freeze the header row: everything from A2 down scrolls under a fixed row 1.
    ws.freeze_panes = "A2"

    wb.save(path)
    return path


def _autosize_columns(ws, posts: list[StoredPost]) -> None:
    """Size each column to its widest cell, clamped to a readable range."""
    for col_idx, (header, attr) in enumerate(COLUMNS, start=1):
        longest = len(header)
        for post in posts:
            value = getattr(post, attr)
            # Measure the longest single line so multi-line post text doesn't
            # blow the width out to the full paragraph length.
            for line in str(value).splitlines() or [""]:
                longest = max(longest, len(line))
        width = max(MIN_COL_WIDTH, min(MAX_COL_WIDTH, longest + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Let the post-text column wrap so long posts stay readable.
    text_col = get_column_letter(_attr_index("text"))
    for row in ws[f"{text_col}2:{text_col}{ws.max_row}"]:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _attr_index(attr: str) -> int:
    for idx, (_header, a) in enumerate(COLUMNS, start=1):
        if a == attr:
            return idx
    raise KeyError(attr)
