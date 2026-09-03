"""Tests for the public (Vercel) upload interface in api/index.py."""

from __future__ import annotations

import io
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "api"))

from api.index import app  # noqa: E402

FIXTURE = ROOT / "tests" / "fixtures" / "sample_search.html"


def _client():
    app.config.update(TESTING=True)
    return app.test_client()


def test_form_page_loads():
    resp = _client().get("/")
    assert resp.status_code == 200
    assert b"Search term" in resp.data
    assert b"Upload the saved HTML file" in resp.data
    # States the no-credentials property to the visitor.
    assert b"stores no credentials" in resp.data


def test_missing_term_is_rejected():
    resp = _client().post("/generate", data={"term": "", "html_text": "<html></html>"})
    assert resp.status_code == 400
    assert b"Please enter a search term." in resp.data


def test_missing_html_is_rejected():
    resp = _client().post("/generate", data={"term": "data engineering"})
    assert resp.status_code == 400
    assert b"upload the saved HTML file" in resp.data


def test_html_with_no_posts_is_rejected():
    resp = _client().post(
        "/generate",
        data={"term": "data engineering", "html_text": "<html><body>nope</body></html>"},
    )
    assert resp.status_code == 400
    assert b"No posts found in that HTML." in resp.data


def test_upload_returns_formatted_xlsx():
    """The whole point: upload saved HTML, get a correct spreadsheet back."""
    data = {
        "term": "data engineering",
        "html_file": (io.BytesIO(FIXTURE.read_bytes()), "results.html"),
    }
    resp = _client().post(
        "/generate", data=data, content_type="multipart/form-data"
    )
    assert resp.status_code == 200
    assert resp.headers["Content-Disposition"].startswith("attachment")
    assert "linkedin_posts_data_engineering.xlsx" in resp.headers[
        "Content-Disposition"
    ]

    ws = load_workbook(io.BytesIO(resp.data)).active
    # Header row, bold and frozen.
    assert [c.value for c in ws[1]] == [
        "Post Link",
        "Date Posted",
        "Posted By",
        "Post Text",
        "Search Term",
        "Captured At",
    ]
    assert ws["A1"].font.bold is True
    assert ws.freeze_panes == "A2"
    # 3 unique posts from the fixture's 4 parsed cards (one duplicate URN).
    assert ws.max_row == 4
    assert ws["C2"].value == "Ada Lovelace"
    assert ws["A2"].hyperlink.target.startswith("https://www.linkedin.com/")
    assert ws["E2"].value == "data engineering"


def test_pasted_html_also_works():
    resp = _client().post(
        "/generate",
        data={"term": "python", "html_text": FIXTURE.read_text()},
    )
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.data)).active
    assert ws.max_row == 4
    assert ws["E2"].value == "python"


if __name__ == "__main__":
    import subprocess

    raise SystemExit(
        subprocess.call([sys.executable, "-m", "pytest", "-q", __file__])
    )
