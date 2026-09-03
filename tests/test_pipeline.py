"""Tests for the pieces that don't need a browser: parse, store, excel, and
the full pipeline wired to a fixture HTML file instead of LinkedIn.

Run with:  python -m pytest -q   (or: python tests/test_pipeline.py)
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from linkedin_posts.cli import run
from linkedin_posts.config import Config
from linkedin_posts.excel import write_xlsx
from linkedin_posts.parse import parse_posts
from linkedin_posts.store import Store

FIXTURE = Path(__file__).parent / "fixtures" / "sample_search.html"


def _fixture_html() -> str:
    return FIXTURE.read_text()


def test_parse_extracts_expected_posts():
    posts = parse_posts(_fixture_html())
    # 5 cards, but 1 is a duplicate URN and 1 has no link. parse_posts keeps
    # both duplicate cards (dedup happens in the store) but skips the linkless.
    assert len(posts) == 4
    first = posts[0]
    assert first.url == (
        "https://www.linkedin.com/feed/update/"
        "urn:li:activity:7200000000000000001/"
    )
    assert first.author == "Ada Lovelace"
    assert first.date_posted == "2d"  # trailing bullet cleaned off
    assert "analytics engine" in first.text
    # The linkless promo card is gone.
    assert all("no linkable identity" not in p.text for p in posts)


def test_store_dedupes_by_url(tmp_path):
    db = tmp_path / "test.db"
    posts = parse_posts(_fixture_html())
    with Store(db) as store:
        result = store.add_posts(posts, "data engineering")
        # 3 unique URLs among the 4 parsed (one duplicate URN).
        assert result.new_count == 3
        assert result.duplicate_count == 1
        assert result.total_count == 3

        # Re-running the same batch adds nothing new (safe re-run).
        again = store.add_posts(posts, "data engineering")
        assert again.new_count == 0
        assert again.duplicate_count == 4
        assert again.total_count == 3


def test_excel_formatting(tmp_path):
    db = tmp_path / "test.db"
    xlsx = tmp_path / "out.xlsx"
    posts = parse_posts(_fixture_html())
    with Store(db) as store:
        store.add_posts(posts, "data engineering")
        write_xlsx(store.all_posts(), xlsx)

    wb = load_workbook(xlsx)
    ws = wb.active
    # Header row is present and bold.
    assert ws["A1"].value == "Post Link"
    assert ws["C1"].value == "Posted By"
    assert ws["D1"].value == "Post Text"
    assert ws["A1"].font.bold is True
    # Header row frozen.
    assert ws.freeze_panes == "A2"
    # Post-link cells carry a real hyperlink.
    assert ws["A2"].hyperlink is not None
    assert ws["A2"].hyperlink.target.startswith("https://www.linkedin.com/")
    # Columns were given an explicit (auto-sized) width.
    assert ws.column_dimensions["A"].width > 0


def test_full_pipeline_with_fixture(tmp_path):
    db = tmp_path / "pipe.db"
    xlsx = tmp_path / "pipe.xlsx"
    config = Config(db_path=str(db), xlsx_path=str(xlsx))

    def fake_provider(term: str, cfg: Config) -> str:
        return _fixture_html()

    code = run("data engineering", config, fake_provider)
    assert code == 0
    assert Path(xlsx).exists()

    # Second run appends nothing (idempotent on the same data).
    code2 = run("data engineering", config, fake_provider)
    assert code2 == 0
    with Store(db) as store:
        assert store.count() == 3


if __name__ == "__main__":
    import subprocess
    import sys

    raise SystemExit(
        subprocess.call([sys.executable, "-m", "pytest", "-q", __file__])
    )
